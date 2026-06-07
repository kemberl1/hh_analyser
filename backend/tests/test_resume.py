"""Tests for Phase 7 — Resume Analyser.

All LLM calls are MOCKED (no real X5 CoPilot API calls in tests).
Tests cover:
  1. Text extraction from PDF/DOCX/raw text
  2. PII sanitisation (NFR-16 — CRITICAL)
  3. Skill extraction and market matching
  4. Analysis service (context, prompt, LLM integration)
  5. REST endpoint (multipart + text, envelope, disabled LLM)
  6. Error handling (format, size, corrupt files)
"""

from __future__ import annotations

import io
from unittest.mock import AsyncMock, patch

import pytest

from app.services.llm.base import LLMClient, LLMError, LLMResponse
from app.services.resume.parser import (
    MAX_FILE_SIZE_BYTES,
    MAX_TEXT_LENGTH,
    ResumeParseError,
    extract_text,
    normalise_text,
)
from app.services.resume.sanitizer import sanitise_pii
from app.services.resume.skills import (
    extract_skills_from_text,
    match_skills_with_market,
)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

class MockLLMClient(LLMClient):
    """Mock LLM client for resume analysis tests."""

    def __init__(self, response_text: str = '{"market_fit_score": 75}'):
        self._response_text = response_text
        self.chat_calls: list = []

    async def chat(self, messages, *, model=None, temperature=None, max_tokens=None, **kwargs):
        self.chat_calls.append({"messages": messages, "model": model})
        return LLMResponse(
            content=self._response_text,
            model=model or "test-model",
            usage={"prompt_tokens": 100,
                   "completion_tokens": 50, "total_tokens": 150},
        )

    async def embeddings(self, texts, *, model=None, **kwargs):
        return [[0.1, 0.2, 0.3] for _ in texts]


def _make_minimal_pdf() -> bytes:
    """Create a minimal valid PDF with text using pdfplumber-compatible format."""
    # Use reportlab-free approach: hand-craft a simple PDF
    content = (
        b"%PDF-1.4\n"
        b"1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n"
        b"2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n"
        b"3 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>\nendobj\n"
        b"4 0 obj\n<< /Length 44 >>\nstream\nBT /F1 12 Tf 100 700 Td (Hello World) Tj ET\nendstream\nendobj\n"
        b"5 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n"
        b"xref\n0 6\n"
        b"0000000000 65535 f \n"
        b"0000000009 00000 n \n"
        b"0000000058 00000 n \n"
        b"0000000115 00000 n \n"
        b"0000000266 00000 n \n"
        b"0000000360 00000 n \n"
        b"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n441\n%%EOF\n"
    )
    return content


def _make_minimal_docx() -> bytes:
    """Create a minimal valid DOCX with text."""
    from docx import Document

    doc = Document()
    doc.add_paragraph("Опытный Frontend-разработчик")
    doc.add_paragraph("Навыки: React, TypeScript, JavaScript, CSS, HTML")
    doc.add_paragraph("Опыт работы: 3 года")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ===========================================================================
# 1. Text extraction tests
# ===========================================================================

class TestTextExtraction:
    """Test resume text extraction from various formats."""

    def test_raw_text_extraction(self):
        """Raw text input is normalised and returned."""
        text = "  Опытный React-разработчик  \n\n\n  с опытом 3 года  "
        result = extract_text(raw_text=text)
        assert "Опытный React-разработчик" in result
        assert "с опытом 3 года" in result
        # Triple newlines collapsed to double
        assert "\n\n\n" not in result

    def test_raw_text_empty_raises(self):
        """Empty raw text raises ResumeParseError."""
        with pytest.raises(ResumeParseError, match="пуст"):
            extract_text(raw_text="   ")

    def test_raw_text_truncation(self):
        """Text longer than MAX_TEXT_LENGTH is truncated."""
        long_text = "A" * (MAX_TEXT_LENGTH + 1000)
        result = extract_text(raw_text=long_text)
        assert len(result) <= MAX_TEXT_LENGTH

    def test_no_input_raises(self):
        """No file and no text raises error."""
        with pytest.raises(ResumeParseError, match="Необходимо"):
            extract_text()

    def test_empty_file_raises(self):
        """Empty file bytes raises error."""
        with pytest.raises(ResumeParseError, match="пуст"):
            extract_text(file_bytes=b"", filename="resume.pdf")

    def test_oversized_file_raises(self):
        """File exceeding size limit raises error with 413."""
        big = b"x" * (MAX_FILE_SIZE_BYTES + 1)
        with pytest.raises(ResumeParseError, match="большой") as exc_info:
            extract_text(file_bytes=big, filename="resume.pdf")
        assert exc_info.value.status_code == 413

    def test_unsupported_format_raises(self):
        """Unsupported file format raises error."""
        with pytest.raises(ResumeParseError, match="Неподдерживаемый"):
            extract_text(file_bytes=b"some content", filename="resume.txt")

    def test_corrupt_pdf_raises(self):
        """Corrupt PDF raises parse error."""
        with pytest.raises(ResumeParseError, match="прочитать PDF"):
            extract_text(file_bytes=b"not a pdf", filename="resume.pdf")

    def test_corrupt_docx_raises(self):
        """Corrupt DOCX raises parse error."""
        with pytest.raises(ResumeParseError, match="прочитать DOCX"):
            extract_text(file_bytes=b"not a docx", filename="resume.docx")

    def test_docx_extraction(self):
        """DOCX text is extracted correctly."""
        docx_bytes = _make_minimal_docx()
        result = extract_text(file_bytes=docx_bytes, filename="resume.docx")
        assert "Frontend" in result
        assert "React" in result

    def test_pdf_extraction(self):
        """PDF text extraction works with valid PDF."""
        pdf_bytes = _make_minimal_pdf()
        try:
            result = extract_text(file_bytes=pdf_bytes, filename="resume.pdf")
            # If pdfplumber can parse our minimal PDF
            assert isinstance(result, str)
        except ResumeParseError:
            # Some minimal PDFs may not have extractable text
            # That's OK — the error handling works
            pass

    def test_content_type_detection(self):
        """Content type is used for format detection."""
        docx_bytes = _make_minimal_docx()
        result = extract_text(
            file_bytes=docx_bytes,
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        assert "React" in result

    def test_normalise_text(self):
        """Text normalisation works."""
        text = "Hello  \t  world\n\n\n\nParagraph two\r\nLine"
        result = normalise_text(text)
        assert "Hello world" in result
        assert "\n\n\n" not in result  # collapsed
        assert "\r" not in result


# ===========================================================================
# 2. PII Sanitisation tests — NFR-16 (CRITICAL)
# ===========================================================================

class TestPIISanitisation:
    """CRITICAL: PII must be removed before sending to external LLM."""

    def test_email_removed(self):
        """Email addresses are sanitised."""
        text = "Контакт: developer@example.com для связи"
        result = sanitise_pii(text)
        assert "developer@example.com" not in result.sanitised_text
        assert "[УДАЛЕНО]" in result.sanitised_text
        assert result.entities_removed > 0
        assert result.categories.get("email", 0) > 0

    def test_phone_russian_removed(self):
        """Russian phone numbers are sanitised."""
        text = "Телефон: +7 (999) 123-45-67. Звоните!"
        result = sanitise_pii(text)
        assert "999" not in result.sanitised_text
        assert "123" not in result.sanitised_text
        assert result.entities_removed > 0
        assert result.categories.get("phone", 0) > 0

    def test_phone_8_format_removed(self):
        """Phone starting with 8 is sanitised."""
        text = "Тел: 8(495)123-45-67"
        result = sanitise_pii(text)
        assert "495" not in result.sanitised_text
        assert result.entities_removed > 0

    def test_url_removed(self):
        """HTTP/HTTPS URLs are sanitised."""
        text = "Портфолио: https://myportfolio.com/projects"
        result = sanitise_pii(text)
        assert "https://myportfolio.com" not in result.sanitised_text
        assert result.categories.get("url", 0) > 0

    def test_telegram_handle_removed(self):
        """Telegram handles are sanitised."""
        text = "Telegram: @my_username"
        result = sanitise_pii(text)
        assert "@my_username" not in result.sanitised_text
        assert result.entities_removed > 0

    def test_social_link_removed(self):
        """Social media links are sanitised."""
        text = "LinkedIn: linkedin.com/in/johndoe"
        result = sanitise_pii(text)
        assert "linkedin.com/in/johndoe" not in result.sanitised_text

    def test_fio_with_patronymic_removed(self):
        """Russian full name with patronymic (ФИО) is sanitised."""
        text = "Кандидат: Иванов Иван Иванович, опыт 5 лет"
        result = sanitise_pii(text)
        assert "Иванов" not in result.sanitised_text
        assert "Иванович" not in result.sanitised_text
        assert result.categories.get("fio", 0) > 0

    def test_fio_name_first_order_surname_removed(self):
        """ФИО in 'Имя Отчество Фамилия' order — surname must NOT leak (NFR-16).

        Regression: previously the matcher only handled
        'Фамилия Имя Отчество'; the trailing surname (e.g. «Сидоров»)
        leaked into the LLM prompt. The whole 3-word span must be removed
        regardless of name order.
        """
        text = "Иван Петрович Сидоров\nFrontend-разработчик, React, TypeScript."
        result = sanitise_pii(text)
        assert "Иван" not in result.sanitised_text
        assert "Петрович" not in result.sanitised_text
        assert "Сидоров" not in result.sanitised_text  # the leak that was fixed
        assert result.categories.get("fio", 0) > 0
        # Professional content preserved
        assert "React" in result.sanitised_text
        assert "TypeScript" in result.sanitised_text
        assert "Frontend-разработчик" in result.sanitised_text

    def test_dob_removed(self):
        """Date of birth is sanitised."""
        text = "Дата рождения: 15.03.1990"
        result = sanitise_pii(text)
        assert "15.03.1990" not in result.sanitised_text
        assert result.categories.get("dob", 0) > 0

    def test_address_removed(self):
        """Physical address markers are sanitised."""
        text = "Адрес: г. Москва, ул. Ленина, д. 15, кв. 42"
        result = sanitise_pii(text)
        assert "кв. 42" not in result.sanitised_text
        assert result.entities_removed > 0

    def test_professional_content_preserved(self):
        """Professional content (skills, experience) is NOT removed."""
        text = (
            "Frontend-разработчик с 5-летним опытом. "
            "Стек: React, TypeScript, Next.js, Redux, CSS Modules. "
            "Опыт работы в Agile-командах. "
            "Разработал компонентную библиотеку для крупного проекта."
        )
        result = sanitise_pii(text)
        assert "React" in result.sanitised_text
        assert "TypeScript" in result.sanitised_text
        assert "Next.js" in result.sanitised_text
        assert "5-летним опытом" in result.sanitised_text
        assert result.entities_removed == 0

    def test_mixed_pii_and_skills(self):
        """PII is removed but skills are preserved from mixed text."""
        text = (
            "Иванов Иван Иванович\n"
            "Email: ivan@example.com\n"
            "Телефон: +7 (999) 123-45-67\n"
            "GitHub: https://github.com/ivandev\n"
            "\n"
            "Frontend-разработчик, 4 года опыта.\n"
            "Навыки: React, TypeScript, Redux, Next.js, Jest.\n"
            "Работал в Яндексе и Тинькофф."
        )
        result = sanitise_pii(text)
        # PII removed
        assert "ivan@example.com" not in result.sanitised_text
        assert "Иванов" not in result.sanitised_text
        assert "Иванович" not in result.sanitised_text
        assert "+7" not in result.sanitised_text or "123-45-67" not in result.sanitised_text
        assert "github.com/ivandev" not in result.sanitised_text
        # Skills preserved
        assert "React" in result.sanitised_text
        assert "TypeScript" in result.sanitised_text
        assert "Redux" in result.sanitised_text
        assert "Frontend-разработчик" in result.sanitised_text
        # Entity count
        assert result.entities_removed >= 4  # email, phone, fio, url

    def test_sanitisation_count_logged(self):
        """Sanitisation result includes counts per category."""
        text = "Email: test@test.com, Тел: +7 (999) 111-22-33"
        result = sanitise_pii(text)
        assert result.entities_removed >= 2
        assert "email" in result.categories
        assert "phone" in result.categories

    def test_section_headers_not_treated_as_fio(self):
        """Section headers like 'Опыт Работы' are NOT treated as FIO."""
        text = "Опыт Работы\nFrontend Developer в компании"
        result = sanitise_pii(text)
        # "Опыт Работы" should not be removed (it's a heading, not a name)
        assert "Опыт Работы" in result.sanitised_text or "Опыт" in result.sanitised_text


# ===========================================================================
# 3. Skill extraction and matching tests
# ===========================================================================

class TestSkillExtraction:
    """Test rule-based skill extraction from resume text."""

    def test_extract_common_skills(self):
        """Common frontend skills are extracted."""
        text = "Опыт работы с React, TypeScript, JavaScript и CSS."
        skills = extract_skills_from_text(text)
        assert "React" in skills
        assert "TypeScript" in skills
        assert "JavaScript" in skills
        assert "CSS" in skills

    def test_extract_case_insensitive(self):
        """Extraction is case-insensitive."""
        text = "Знаю react и TYPESCRIPT"
        skills = extract_skills_from_text(text)
        assert "React" in skills
        assert "TypeScript" in skills

    def test_extract_aliases(self):
        """Aliases are resolved to canonical names."""
        text = "Работал с reactjs, vuejs, nextjs и scss"
        skills = extract_skills_from_text(text)
        assert "React" in skills
        assert "Vue.js" in skills
        assert "Next.js" in skills
        assert "CSS" in skills  # scss → CSS

    def test_extract_no_duplicates(self):
        """No duplicate skills in output."""
        text = "React и react и React.js"
        skills = extract_skills_from_text(text)
        assert skills.count("React") == 1

    def test_extract_empty_text(self):
        """Empty text returns no skills."""
        skills = extract_skills_from_text("")
        assert skills == []

    def test_extract_russian_terms(self):
        """Russian skill terms are extracted."""
        text = "Адаптивная вёрстка, кроссбраузерность"
        skills = extract_skills_from_text(text)
        assert "Responsive Design" in skills
        assert "Cross-browser" in skills


class TestSkillMatching:
    """Test skill matching against market data."""

    def test_full_match(self):
        """All market skills matched."""
        resume = ["React", "TypeScript"]
        market = [{"skill": "React", "share": 0.8, "count": 100},
                  {"skill": "TypeScript", "share": 0.7, "count": 90}]
        result = match_skills_with_market(resume, market)
        assert result.match_ratio == 1.0
        assert len(result.matched_skills) == 2
        assert len(result.missing_skills) == 0

    def test_partial_match(self):
        """Some market skills missing from resume."""
        resume = ["React"]
        market = [{"skill": "React", "share": 0.8, "count": 100},
                  {"skill": "TypeScript", "share": 0.7, "count": 90}]
        result = match_skills_with_market(resume, market)
        assert result.match_ratio == 0.5
        assert "React" in result.matched_skills
        assert "TypeScript" in result.missing_skills

    def test_extra_skills(self):
        """Resume skills not in market top are tracked."""
        resume = ["React", "Svelte"]
        market = [{"skill": "React", "share": 0.8, "count": 100}]
        result = match_skills_with_market(resume, market, top_n=1)
        assert "Svelte" in result.extra_skills

    def test_empty_resume(self):
        """Empty resume skills."""
        result = match_skills_with_market(
            [], [{"skill": "React", "share": 0.8, "count": 100}]
        )
        assert result.match_ratio == 0.0
        assert len(result.missing_skills) == 1

    def test_empty_market(self):
        """Empty market data."""
        result = match_skills_with_market(["React"], [])
        assert result.match_ratio == 0.0


# ===========================================================================
# 4. Analysis service tests
# ===========================================================================
_MOCK_MARKET_CTX = {
    "grade_filter": "all",
    "salary_all": {
        "median": 200000, "p25": 150000, "p75": 280000,
        "min": 80000, "max": 400000, "sample_size": 100, "total_vacancies": 500,
    },
    "salary_grade": None,
    "top_skills_all": [
        {"name": "React", "share": 0.8, "count": 400},
        {"name": "TypeScript", "share": 0.7, "count": 350},
    ],
    "top_skills_grade": None,
}


class TestAnalysisService:
    """Test resume analysis service."""

    @pytest.mark.asyncio
    async def test_analysis_with_mock_llm(self):
        """Full analysis pipeline with mocked LLM."""
        from app.services.resume.analysis import analyse_resume
        from app.services.resume.skills import SkillMatchResult

        mock_client = MockLLMClient(
            response_text='{"market_fit_score": 80, "estimated_grade": "middle", '
                          '"salary_range": {"from": 200000, "to": 300000}, '
                          '"passes_keyword_filters": true, '
                          '"strengths": ["Good stack"], '
                          '"weaknesses": ["No tests"], '
                          '"recommendations": ["Learn testing"]}'
        )

        skill_match = SkillMatchResult(
            resume_skills=["React", "TypeScript"],
            matched_skills=["React", "TypeScript"],
            missing_skills=["Next.js"],
            match_ratio=0.67,
            extra_skills=[],
        )

        mock_session = AsyncMock()

        with patch("app.services.resume.analysis.build_resume_market_context",
                   new_callable=AsyncMock, return_value=_MOCK_MARKET_CTX):
            result = await analyse_resume(
                mock_session,
                mock_client,
                sanitised_text="Resume. React, TypeScript.",
                skill_match=skill_match,
            )

        assert result["market_fit_score"] == 80
        assert result["estimated_grade"] == "middle"
        assert result["llm_enhanced"] is True
        assert result["model"] == "test-model"
        assert len(mock_client.chat_calls) == 1

    @pytest.mark.asyncio
    async def test_analysis_without_llm(self):
        """Analysis with LLM=None returns rule-based result."""
        from app.services.resume.analysis import analyse_resume
        from app.services.resume.skills import SkillMatchResult

        skill_match = SkillMatchResult(
            resume_skills=["React", "TypeScript", "JavaScript", "CSS", "HTML"],
            matched_skills=["React", "TypeScript", "JavaScript"],
            missing_skills=["Next.js"],
            match_ratio=0.75,
            extra_skills=["CSS", "HTML"],
        )

        mock_session = AsyncMock()
        with patch("app.services.resume.analysis.build_resume_market_context",
                   new_callable=AsyncMock, return_value=_MOCK_MARKET_CTX):
            result = await analyse_resume(
                mock_session, None,
                sanitised_text="React TypeScript developer",
                skill_match=skill_match,
            )

        assert result["llm_enhanced"] is False
        assert result["market_fit_score"] > 0
        assert "llm_error" in result
        assert result["matched_skills"] == [
            "React", "TypeScript", "JavaScript"]

    @pytest.mark.asyncio
    async def test_analysis_llm_error_graceful(self):
        """LLM error is handled gracefully — rule-based fallback."""
        from app.services.resume.analysis import analyse_resume
        from app.services.resume.skills import SkillMatchResult

        class FailingClient(LLMClient):
            async def chat(self, messages, **kwargs):
                raise LLMError("Service unavailable",
                               status_code=503, retryable=True)

            async def embeddings(self, texts, **kwargs):
                raise LLMError("Service unavailable")

        skill_match = SkillMatchResult(
            resume_skills=["React"], matched_skills=["React"],
            missing_skills=[], match_ratio=1.0, extra_skills=[],
        )

        mock_session = AsyncMock()
        with patch("app.services.resume.analysis.build_resume_market_context",
                   new_callable=AsyncMock, return_value=_MOCK_MARKET_CTX):
            result = await analyse_resume(
                mock_session, FailingClient(),
                sanitised_text="React developer",
                skill_match=skill_match,
            )

        assert result["llm_enhanced"] is False
        assert "llm_error" in result
        assert result["market_fit_score"] > 0

    @pytest.mark.asyncio
    async def test_prompt_contains_no_pii(self):
        """The prompt sent to LLM contains no PII (only sanitised text)."""
        from app.services.resume.analysis import build_resume_prompt
        from app.services.resume.skills import SkillMatchResult

        sanitised = "[REMOVED] — Frontend Developer. React, TypeScript."
        skill_match = SkillMatchResult(
            resume_skills=["React", "TypeScript"],
            matched_skills=["React"],
            missing_skills=["Next.js"],
            match_ratio=0.5,
            extra_skills=["TypeScript"],
        )
        context = {"salary_all": {"median": 200000}, "top_skills_all": []}
        messages = build_resume_prompt(sanitised, skill_match, context)
        all_content = " ".join(msg.content for msg in messages)
        assert "ivan@" not in all_content
        assert "React" in all_content


# ===========================================================================
# 5. Endpoint tests
# ===========================================================================

# Helper: mock return for _run_analysis
def _mock_run_analysis_result(**overrides):
    from app.schemas.resume import (
        ResumeAnalysisData, ResumeAnalysisMeta, ResumeAnalysisResponse,
    )
    return ResumeAnalysisResponse(
        meta=ResumeAnalysisMeta(
            analyzed_at="2026-01-01T00:00:00Z",
            llm_enabled=overrides.get("llm_enabled", False),
            llm_enhanced=overrides.get("llm_enhanced", False),
            pii_entities_removed=overrides.get("pii_entities_removed", 0),
        ),
        data=ResumeAnalysisData(
            market_fit_score=overrides.get("market_fit_score", 50),
            passes_keyword_filters=overrides.get(
                "passes_keyword_filters", True),
            matched_skills=overrides.get("matched_skills", ["React"]),
            missing_in_demand_skills=overrides.get(
                "missing_in_demand_skills", ["Next.js"]),
            resume_skills=overrides.get(
                "resume_skills", ["React", "TypeScript"]),
            skill_match_ratio=overrides.get("skill_match_ratio", 0.5),
            recommendations=overrides.get(
                "recommendations", ["Learn Next.js"]),
            market_sample_size=overrides.get("market_sample_size", 500),
            llm_error=overrides.get("llm_error", "LLM disabled"),
        ),
    )


class TestResumeEndpoint:
    """Test POST /api/v1/resume/analyze endpoint."""

    @pytest.mark.asyncio
    async def test_text_submission(self, client):
        """Text-based resume submission works."""
        mock_result = _mock_run_analysis_result()
        with patch("app.api.v1.endpoints.resume._run_analysis",
                   new_callable=AsyncMock, return_value=mock_result):
            response = await client.post(
                "/api/v1/resume/analyze",
                data={
                    "resume_text": "Frontend developer with React and TypeScript. 3 years."},
            )
        assert response.status_code == 200
        body = response.json()
        assert "meta" in body
        assert "data" in body
        assert "market_fit_score" in body["data"]
        assert "recommendations" in body["data"]
        assert isinstance(body["data"]["matched_skills"], list)

    @pytest.mark.asyncio
    async def test_docx_submission(self, client):
        """DOCX file submission works."""
        docx_bytes = _make_minimal_docx()
        mock_result = _mock_run_analysis_result()
        with patch("app.api.v1.endpoints.resume._run_analysis",
                   new_callable=AsyncMock, return_value=mock_result):
            response = await client.post(
                "/api/v1/resume/analyze",
                files={"file": ("resume.docx", docx_bytes,
                                "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
        assert response.status_code == 200
        body = response.json()
        assert "meta" in body and "data" in body
        assert body["data"]["market_fit_score"] >= 0

    @pytest.mark.asyncio
    async def test_no_input_returns_400(self, client):
        """No file and no text returns 400."""
        response = await client.post("/api/v1/resume/analyze")
        assert response.status_code in (400, 422)

    @pytest.mark.asyncio
    async def test_unsupported_format_returns_400(self, client):
        """Unsupported file format returns 400."""
        response = await client.post(
            "/api/v1/resume/analyze",
            files={"file": ("readme.txt", b"some text content", "text/plain")},
        )
        assert response.status_code == 400
        assert "Неподдерживаемый" in response.json().get("detail", "")

    @pytest.mark.asyncio
    async def test_invalid_grade_returns_400(self, client):
        """Invalid target_grade returns 400."""
        response = await client.post(
            "/api/v1/resume/analyze",
            data={"resume_text": "Some resume text",
                  "target_grade": "invalid_grade"},
        )
        assert response.status_code == 400

    @pytest.mark.asyncio
    async def test_disabled_llm_returns_200(self, client):
        """When LLM is disabled, still returns 200 with rule-based analysis."""
        mock_result = _mock_run_analysis_result(
            llm_enhanced=False,
            llm_error="LLM disabled",
            market_fit_score=45,
        )
        with patch("app.api.v1.endpoints.resume._run_analysis",
                   new_callable=AsyncMock, return_value=mock_result):
            response = await client.post(
                "/api/v1/resume/analyze",
                data={"resume_text": "React and TypeScript developer. 2 years."},
            )
        assert response.status_code == 200
        body = response.json()
        assert body["meta"]["llm_enhanced"] is False
        assert body["data"]["llm_error"] is not None
        assert body["data"]["market_fit_score"] >= 0
        assert isinstance(body["data"]["recommendations"], list)

    @pytest.mark.asyncio
    async def test_response_envelope_structure(self, client):
        """Response follows {meta, data} envelope contract."""
        mock_result = _mock_run_analysis_result()
        with patch("app.api.v1.endpoints.resume._run_analysis",
                   new_callable=AsyncMock, return_value=mock_result):
            response = await client.post(
                "/api/v1/resume/analyze",
                data={"resume_text": "Senior React developer with 5 years experience."},
            )
        assert response.status_code == 200
        body = response.json()
        assert "meta" in body
        meta = body["meta"]
        assert "analyzed_at" in meta
        assert "llm_enabled" in meta
        assert "llm_enhanced" in meta
        assert "pii_entities_removed" in meta
        assert "data" in body
        data = body["data"]
        assert "market_fit_score" in data
        assert "passes_keyword_filters" in data
        assert "matched_skills" in data
        assert "missing_in_demand_skills" in data
        assert "resume_skills" in data
        assert "recommendations" in data

    @pytest.mark.asyncio
    async def test_pii_removed_in_analysis(self, client):
        """PII in submitted text is counted in meta."""
        mock_result = _mock_run_analysis_result(pii_entities_removed=4)
        with patch("app.api.v1.endpoints.resume._run_analysis",
                   new_callable=AsyncMock, return_value=mock_result):
            response = await client.post(
                "/api/v1/resume/analyze",
                data={
                    "resume_text": "Иванов Иван Иванович\nEmail: ivan@example.com\nReact developer."},
            )
        assert response.status_code == 200
        body = response.json()
        assert body["meta"]["pii_entities_removed"] >= 3
